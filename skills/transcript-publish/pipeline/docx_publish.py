"""Render an accessible .docx transcript from cues + a PublishJudgment.

Mirrors publish.py's build_html/build_body exactly in structure and rules
(same section/paragraph/speaker/non-speech logic, same doubled-word fix, same
"never add text" constraint) but emits a python-docx Document instead of an
HTML string. Kept as a sibling module — rather than folded into publish.py —
so an HTML-only install never needs python-docx as a dependency.
"""

from __future__ import annotations

from docx import Document
from docx.shared import Pt, RGBColor

from .parse import Cue, timecode_to_seconds
from .judgment import PublishJudgment
from .vtt import duration_minutes

_GRAY = RGBColor(0x55, 0x55, 0x55)


def _mmss(vtt_timecode: str) -> str:
    secs = timecode_to_seconds(vtt_timecode)
    h = int(secs // 3600)
    m = int((secs % 3600) // 60)
    s = int(secs % 60)
    return f"{h}:{m:02d}:{s:02d}" if h else f"{m:02d}:{s:02d}"


# ---- Doubled-word cleanup at publish (identical rule to publish.py) -------

def _apply_publish_text_fixes(cue: Cue, judgment: PublishJudgment) -> str:
    """The only text edit the publish step may make: doubled-word removal.
    See publish.py's docstring of the same name for the full rationale."""
    text = cue.text
    for d in judgment.doubled_words:
        if d.get("entry") == cue.index and d.get("find") in text:
            text = text.replace(d["find"], d.get("replace", d["find"]), 1)
    return text


_MIN_PT = Pt(12)  # every run in the document must be at least this size


def _add_timecode_paragraph(doc: Document, vtt_start: str) -> None:
    p = doc.add_paragraph()
    run = p.add_run(f"[{_mmss(vtt_start)}]")
    run.font.size = _MIN_PT
    run.font.color.rgb = _GRAY
    p.paragraph_format.space_before = Pt(12)
    p.paragraph_format.space_after = Pt(2)


def _add_non_speech_paragraph(doc: Document, description: str) -> None:
    p = doc.add_paragraph()
    run = p.add_run(f"[{description}]")
    run.italic = True
    run.font.color.rgb = _GRAY


def build_body(doc: Document, cues: list[Cue], judgment: PublishJudgment) -> None:
    section_starts = {s["start_entry"]: s for s in judgment.sections}
    breaks = set(judgment.paragraph_breaks)
    speaker_at = {s["entry"]: s.get("speaker", "")
                  for s in judgment.speaker_turns if s.get("speaker")}
    non_speech_at = {n["entry"]: n.get("description", "")
                     for n in judgment.non_speech if n.get("description")}
    first_section_emitted = False

    para_words: list[str] = []
    para_speaker: str | None = None

    def flush_para():
        nonlocal para_speaker
        if para_words:
            p = doc.add_paragraph()
            if para_speaker:
                label_run = p.add_run(f"{para_speaker}: ")
                label_run.bold = True
            p.add_run(" ".join(para_words))
            para_words.clear()
        para_speaker = None

    for cue in cues:
        if cue.index in section_starts:
            flush_para()
            sec = section_starts[cue.index]
            level = sec.get("level", 2)
            if not (level == 2 and not first_section_emitted):
                _add_timecode_paragraph(doc, cue.start)
            # Bracketed: this heading is an AI-derived topic label, not
            # something the speaker said (same convention as [Applause] and
            # [unintelligible] elsewhere in this pipeline).
            doc.add_heading(f"[{sec['title']}]", level=level)
            if level == 2:
                first_section_emitted = True

        if cue.index in non_speech_at:
            flush_para()
            _add_non_speech_paragraph(doc, non_speech_at[cue.index])
            continue

        if cue.index in speaker_at:
            flush_para()
            para_speaker = speaker_at[cue.index]
        elif cue.index in breaks and para_words:
            flush_para()

        text = _apply_publish_text_fixes(cue, judgment)
        para_words.append(text)

    flush_para()


def _set_minimum_font_sizes(doc: Document) -> None:
    """Every style used in this document must render at >= 12pt. Word's
    default template has Normal at ~11pt and Heading 3 unset (inherits an
    11pt docDefault), both below the floor — set every style explicitly
    rather than trust template defaults."""
    doc.styles["Normal"].font.size = _MIN_PT
    doc.styles["Heading 1"].font.size = Pt(18)
    doc.styles["Heading 2"].font.size = Pt(15)
    doc.styles["Heading 3"].font.size = Pt(13)


def build_docx(cues: list[Cue], judgment: PublishJudgment,
               *, speaker: str = "", course: str = "",
               title: str | None = None) -> Document:
    doc = Document()
    _set_minimum_font_sizes(doc)

    final_title = title or judgment.title or "Transcript"
    # NOT bracketed: unlike the section headers in build_body, the document
    # title is the document's name, not a label inserted into the spoken
    # content — so it doesn't carry the "this wasn't said aloud" marker.
    doc.add_heading(final_title, level=1)

    def _meta_line(label: str, value: str):
        p = doc.add_paragraph()
        r = p.add_run(f"{label}: ")
        r.bold = True
        p.add_run(value)
        return p

    if speaker:
        _meta_line("Speaker", speaker)
    if course:
        _meta_line("Course", course)
    _meta_line("Duration", f"{duration_minutes(cues)} minutes (approx.)")

    note = doc.add_paragraph()
    note_run = note.add_run("Auto-generated transcript. Edits have been applied for clarity.")
    note_run.italic = True

    doc.add_paragraph()  # spacer before body

    build_body(doc, cues, judgment)
    return doc
